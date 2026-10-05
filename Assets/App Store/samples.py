"""Writes the sample documents the App Store screenshots open.

Usage: python3 samples.py <iPhone|iPad> <en|ja> <output directory>

The iPhone set is sized to fit a phone screen without scrolling; the iPad set is
larger so the sheets fill the screen. Everything is set in Kivotos. Needs
openpyxl (pip3 install openpyxl).
"""
import csv, os, random, re, shutil, sys, zipfile, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as L

# The output directory is emptied before it is written, so it is only accepted
# where capture.sh puts it: one folder per device and language under the
# samples root, never anywhere else on disk.
SAMPLES_ROOT = os.path.realpath("/tmp/tables-screenshots-samples")
if len(sys.argv) != 4 or sys.argv[1] not in ("iPhone", "iPad") or sys.argv[2] not in ("en", "ja"):
    sys.exit(__doc__)
DEVICE, LANG = sys.argv[1], sys.argv[2]
OUT = os.path.join(SAMPLES_ROOT, f"{DEVICE}-{LANG}")
if os.path.realpath(sys.argv[3]) != OUT:
    sys.exit(f"The output directory must be {OUT}")
JA = LANG == "ja"
shutil.rmtree(OUT, ignore_errors=True)
os.makedirs(OUT)
rnd = random.Random(7)


def t(en, ja):
    return ja if JA else en


# MARK: - Styles

def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


thin = Side(style="thin", color="BFBFBF")
med = Side(style="medium", color="1F3864")
dbl = Side(style="double", color="1F3864")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
CREDITS = "#,##0;-#,##0"
DAY = "m/d" if JA else "d mmm"
MONTHS = [f"{m}月" for m in range(1, 13)] if JA else \
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
SCHOOL_TINTS = {"Abydos": "FFF2CC", "Millennium": "DDEBF7", "Gehenna": "FBE5D6", "Trinity": "FFF9E6",
                "SRT": "E2EFDA", "Valkyrie": "D9E1F2", "Hyakkiyako": "EDE1F5", "Arius": "E7E6E6",
                "Wildhunt": "F8DDEB", "Odyssey": "D5EEF5"}
SCHOOL_NAMES = {"Abydos": "アビドス", "Millennium": "ミレニアム", "Gehenna": "ゲヘナ", "Trinity": "トリニティ",
                "SRT": "SRT", "Valkyrie": "ヴァルキューレ", "Hyakkiyako": "百鬼夜行", "Arius": "アリウス",
                "Wildhunt": "ワイルドハント", "Odyssey": "オデュッセイア"}
STATUS = {"done": t("Done", "完了"), "doing": t("In Progress", "進行中"),
          "todo": t("Not Started", "未着手"), "blocked": t("Blocked", "保留中")}
STATUS_FILL = {"done": "C6EFCE", "doing": "FFEB9C", "todo": "F2F2F2", "blocked": "FFC7CE"}
STATUS_FONT = {"done": "006100", "doing": "9C5700", "todo": "595959", "blocked": "9C0006"}


def school(key):
    return SCHOOL_NAMES[key] if JA else key


def header(ws, row, cols, color):
    for c in range(1, cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill(color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = box


def grid(ws, r1, r2, c1, c2):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            ws.cell(row=r, column=c).border = box


def widths(ws, ws_widths):
    for col, w in ws_widths.items():
        ws.column_dimensions[col].width = w


def title(ws, text, last_col, size, color, background=None, height=30):
    ws.merge_cells(f"A1:{last_col}1")
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=size, color=color)
    if background:
        ws["A1"].fill = fill(background)
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    else:
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = height


def status_cell(cell, status):
    cell.value = STATUS[status]
    cell.fill = fill(STATUS_FILL[status])
    cell.font = Font(color=STATUS_FONT[status], bold=True)
    cell.alignment = Alignment(horizontal="center")


def placeholder_sheets(wb, *names):
    for name in names:
        wb.create_sheet(name)["A1"] = name


def save(wb, name):
    """Writes the workbook without the parts Tables would flag.

    openpyxl always adds docProps/, which Tables reports as an unsupported
    feature, and writes colours with an alpha of 00. Excel ignores that alpha;
    Tables currently draws it as transparent, so it is raised to FF here.
    """
    path = os.path.join(OUT, name)
    wb.save(path)
    tmp = path + ".tmp"
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            if item.filename.startswith("docProps/"):
                continue
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = re.sub(rb'<Override[^>]*PartName="/docProps/[^"]*"[^>]*/>', b"", data)
            if item.filename == "_rels/.rels":
                data = re.sub(rb'<Relationship[^>]*Target="[^"]*docProps/[^"]*"[^>]*/>', b"", data)
            if item.filename == "xl/styles.xml":
                data = data.replace(b'rgb="00', b'rgb="FF')
            dst.writestr(item, data)
    os.replace(tmp, path)


# MARK: - Foreclosure Task Force budget

EXPENSES = [
    ("Loan repayment", "借金返済", 788000, 0), ("Loan interest", "利息", 96000, 0),
    ("Ammunition", "弾薬", 42000, 0.3), ("Water & power", "水道光熱費", 18000, 0.15),
    ("Ramen", "ラーメン", 9000, 0.4), ("Bicycle parts", "自転車パーツ", 6000, 0.5),
    ("Sand removal", "砂の除去", 15000, 0.2), ("School repairs", "校舎修繕", 30000, 0.6),
    ("Uniforms", "制服", 8000, 0.9), ("Snacks", "おやつ", 3000, 0.4),
    ("Medicine", "医薬品", 5000, 0.4), ("Fuel", "燃料", 7000, 0.2),
    ("Club supplies", "部活用品", 4000, 0.4), ("Stationery", "文房具", 2000, 0.3),
    ("Emergency fund", "緊急資金", 20000, 0), ("Savings", "貯金", 10000, 0),
]


def budget_iphone():
    wb = Workbook()
    ws = wb.active
    ws.title = t("October", "10月")
    title(ws, t("Foreclosure Task Force", "対策委員会 予算"), "E", 18, "1F3864")
    ws.append([])
    ws.append([t("Category", "項目"), t("Budget", "予算"), t("Spent", "支出"), t("Left", "残り"), t("Used", "消化率")])
    header(ws, 3, 5, "1F4E79")
    spent = [788000, 96000, 51800, 16420, 12600, 4380, 15000, 21750, 0, 4150, 3240, 6880, 2970, 1850, 20000, 10000]
    over = []
    for i, ((en, ja, budget, _), s) in enumerate(zip(EXPENSES, spent), start=4):
        ws.append([t(en, ja), budget, s, f"=B{i}-C{i}", f"=C{i}/B{i}"])
        for col in "BCD":
            ws[f"{col}{i}"].number_format = CREDITS
        ws[f"E{i}"].number_format = "0.0%"
        if i % 2 == 0:
            for col in "ABCDE":
                ws[f"{col}{i}"].fill = fill("EAF1FB")
        if s > budget:
            over.append(i)
    last = 3 + len(EXPENSES)
    tr = last + 1
    ws.append([t("Total", "合計"), f"=SUM(B4:B{last})", f"=SUM(C4:C{last})", f"=B{tr}-C{tr}", f"=C{tr}/B{tr}"])
    for col in "ABCDE":
        cell = ws[f"{col}{tr}"]
        cell.font = Font(bold=True)
        cell.border = Border(top=med, bottom=dbl, left=thin, right=thin)
        cell.fill = fill("D9E2F3")
        cell.number_format = "0.0%" if col == "E" else CREDITS
    grid(ws, 4, last, 1, 5)
    for r in over:
        ws[f"D{r}"].font = Font(color="C00000", bold=True)
    widths(ws, {"A": 18, "B": 13, "C": 13, "D": 12, "E": 10})

    inc = wb.create_sheet(t("Income", "収入"))
    inc.append([t("Source", "収入源"), MONTHS[7], MONTHS[8], MONTHS[9], t("Total", "合計")])
    header(inc, 1, 5, "375623")
    for i, (en, ja, *m) in enumerate([("Bounties", "賞金", 520000, 610000, 655000),
                                      ("Part-time jobs", "アルバイト", 118000, 124000, 121500),
                                      ("Scrap sales", "スクラップ売却", 64000, 81000, 77300),
                                      ("Nonomi's card", "ノノミのカード", 300000, 300000, 300000)], start=2):
        inc.append([t(en, ja), *m, f"=SUM(B{i}:D{i})"])
        for col in "BCDE":
            inc[f"{col}{i}"].number_format = CREDITS
    grid(inc, 2, 5, 1, 5)
    widths(inc, {"A": 16, "B": 11, "C": 11, "D": 11, "E": 12})

    debt = wb.create_sheet(t("Debt", "借金"))
    debt.append([t("Lender", "借入先"), t("Balance", "残高"), t("Monthly", "月額"), t("Months left", "残り月数")])
    header(debt, 1, 4, "C00000")
    debt.append([t("Kaiser Loan", "カイザーローン"), 962340000, 788000, "=B2/C2"])
    debt["B2"].number_format = CREDITS
    debt["C2"].number_format = CREDITS
    debt["D2"].number_format = "#,##0"
    grid(debt, 2, 2, 1, 4)
    widths(debt, {"A": 16, "B": 14, "C": 11, "D": 12})
    save(wb, t("Abydos Budget.xlsx", "アビドス予算.xlsx"))


def budget_ipad():
    wb = Workbook()
    ws = wb.active
    ws.title = "2026"
    title(ws, t("Foreclosure Task Force — 2026 Budget", "対策委員会 2026年度 予算"), "N", 18, "1F3864", height=32)

    def section(row, label, color):
        ws.cell(row=row, column=1, value=label)
        for i, m in enumerate(MONTHS, start=2):
            ws.cell(row=row, column=i, value=m)
        ws.cell(row=row, column=14, value=t("Total", "合計"))
        header(ws, row, 14, color)
        ws.cell(row=row, column=1).alignment = Alignment(horizontal="left", vertical="center", indent=1)

    def money_row(row, label, base, jitter, stripe):
        # A tenth of the iPhone figures, so a month's total fits its column.
        base //= 10
        ws.cell(row=row, column=1, value=label)
        for c in range(2, 14):
            v = base if jitter == 0 else round(base * (1 + rnd.uniform(-jitter, jitter)), -1)
            ws.cell(row=row, column=c, value=v).number_format = CREDITS
        ws.cell(row=row, column=14, value=f"=SUM(B{row}:M{row})").number_format = CREDITS
        ws.cell(row=row, column=14).font = Font(bold=True)
        if stripe:
            for c in range(1, 15):
                ws.cell(row=row, column=c).fill = fill(stripe)

    def total_row(row, label, first, last, color):
        ws.cell(row=row, column=1, value=label)
        for c in range(2, 15):
            col = L(c)
            ws.cell(row=row, column=c, value=f"=SUM({col}{first}:{col}{last})").number_format = CREDITS
        for c in range(1, 15):
            cell = ws.cell(row=row, column=c)
            cell.font = Font(bold=True)
            cell.fill = fill(color)
            cell.border = Border(top=med, bottom=dbl, left=thin, right=thin)

    section(3, t("Income", "収入"), "375623")
    income = [("Bounties", "賞金", 600000, 0.25), ("Part-time jobs", "アルバイト", 120000, 0.1),
              ("Scrap sales", "スクラップ売却", 75000, 0.4), ("Nonomi's card", "ノノミのカード", 300000, 0)]
    for i, (en, ja, b, j) in enumerate(income, start=4):
        money_row(i, t(en, ja), b, j, "EEF6EA" if i % 2 == 0 else None)
    grid(ws, 4, 7, 1, 14)
    total_row(8, t("Total income", "収入合計"), 4, 7, "D9EAD3")

    section(10, t("Expenses", "支出"), "1F4E79")
    first = 11
    for i, (en, ja, b, j) in enumerate(EXPENSES, start=first):
        money_row(i, t(en, ja), b, j, "EAF1FB" if i % 2 == 0 else None)
    last = first + len(EXPENSES) - 1
    grid(ws, first, last, 1, 14)
    tx = last + 1
    total_row(tx, t("Total expenses", "支出合計"), first, last, "D9E2F3")

    net = tx + 2
    ws.cell(row=net, column=1, value=t("Net", "収支"))
    ws.cell(row=net + 1, column=1, value=t("Year to date", "累計"))
    for c in range(2, 15):
        col = L(c)
        ws.cell(row=net, column=c, value=f"={col}8-{col}{tx}").number_format = CREDITS
        if c < 14:
            ws.cell(row=net + 1, column=c, value=f"=SUM($B${net}:{col}{net})").number_format = CREDITS
    for r in (net, net + 1):
        for c in range(1, 15):
            ws.cell(row=r, column=c).font = Font(bold=True, color="1F3864")
            ws.cell(row=r, column=c).border = box
            ws.cell(row=r, column=c).fill = fill("FFF2CC")
    widths(ws, {"A": 19, **{L(c): 11 for c in range(2, 14)}, "N": 14})
    placeholder_sheets(wb, t("Income", "収入"), t("Debt", "借金"))
    save(wb, t("Abydos Budget.xlsx", "アビドス予算.xlsx"))


# MARK: - Joint festival plan

def plan_iphone():
    wb = Workbook()
    ws = wb.active
    ws.title = t("Timeline", "日程")
    title(ws, t("Kivotos Joint Festival", "キヴォトス合同祭 計画表"), "F", 16, "FFFFFF", "203864", 28)
    ws.append([t("Task", "タスク"), t("Owner", "担当"), t("Start", "開始"), t("End", "終了"),
               t("Days", "日数"), t("Status", "状況")])
    header(ws, 2, 6, "2F5597")
    tasks = [("Arona briefing", "アロナと相談", "Arona", "アロナ", (9, 1), (9, 2), "done"),
             ("Budget review", "予算承認", "Yuuka", "ユウカ", (9, 3), (9, 9), "done"),
             ("Venue security", "会場警備", "Hina", "ヒナ", (9, 7), (9, 18), "done"),
             ("Food stalls", "屋台の手配", "Fuuka", "フウカ", (9, 14), (10, 9), "doing"),
             ("Game booth", "ゲームブース", "Momoi", "モモイ", (9, 21), (10, 16), "doing"),
             ("Tea party", "お茶会", "Nagisa", "ナギサ", (10, 5), (10, 23), "todo"),
             ("Stage lighting", "舞台照明", "Utaha", "ウタハ", (10, 12), (10, 30), "todo"),
             ("Fireworks", "花火", "Aru", "アル", (10, 19), (11, 6), "blocked"),
             ("Cleanup crew", "片付け班", "Serika", "セリカ", (11, 2), (11, 9), "todo"),
             ("Opening day", "開幕", "Sensei", "先生", (11, 14), (11, 14), "todo")]
    for i, (en, ja, owner_en, owner_ja, s, e, st) in enumerate(tasks, start=3):
        ws.append([t(en, ja), t(owner_en, owner_ja), dt.date(2026, *s), dt.date(2026, *e), f"=D{i}-C{i}+1"])
        ws[f"C{i}"].number_format = DAY
        ws[f"D{i}"].number_format = DAY
        ws[f"E{i}"].alignment = Alignment(horizontal="center")
        status_cell(ws[f"F{i}"], st)
    grid(ws, 3, 2 + len(tasks), 1, 6)
    widths(ws, {"A": 17, "B": 10, "C": 9, "D": 9, "E": 8, "F": 15})

    notes = wb.create_sheet(t("Notes", "メモ"))
    notes.append([t("Date", "日付"), t("Note", "メモ")])
    header(notes, 1, 2, "2F5597")
    for d, en, ja in [((9, 2), "Every school has agreed to take part. Arona will keep the schedule.",
                       "全学園の参加が決定。日程はアロナが管理。"),
                      ((9, 9), "Yuuka approved the budget after cutting the game booth's request in half.",
                       "ゲームブースの申請額を半分にして、ユウカが予算を承認。"),
                      ((10, 20), "The fireworks need a permit before Problem Solver 68 can set them up.",
                       "便利屋68が花火を設置する前に許可が必要。")]:
        notes.append([dt.date(2026, *d), t(en, ja)])
    for r in range(2, 5):
        notes[f"A{r}"].number_format = DAY
        notes[f"A{r}"].alignment = Alignment(vertical="top")
        notes[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
        notes.row_dimensions[r].height = 46
    grid(notes, 2, 4, 1, 2)
    widths(notes, {"A": 10, "B": 40})
    save(wb, t("Festival Plan.xlsx", "合同祭計画.xlsx"))


def plan_ipad():
    wb = Workbook()
    ws = wb.active
    ws.title = t("Timeline", "日程")
    title(ws, t("Kivotos Joint Festival — Plan", "キヴォトス合同祭 計画表"), "H", 18, "FFFFFF", "203864", 34)
    ws.append([t("Task", "タスク"), t("Owner", "担当"), t("Start", "開始"), t("End", "終了"),
               t("Days", "日数"), t("Done", "進捗"), t("Status", "状況"), t("Notes", "メモ")])
    header(ws, 2, 8, "2F5597")
    phases = [
        ("Planning", "企画", "D9E1F2", [
            ("Meet with Arona", "アロナと打ち合わせ", "Arona", "アロナ", (9, 1), (9, 2), 1, "Dates agreed with every school", "全学園と日程を合意"),
            ("Budget approval", "予算承認", "Yuuka", "ユウカ", (9, 3), (9, 9), 1, "Game booth request halved", "ゲームブースの申請額は半額に"),
            ("Joint committee", "合同委員会", "Nagisa", "ナギサ", (9, 3), (9, 11), 1, "", ""),
            ("Master schedule", "全体スケジュール", "Noa", "ノア", (9, 8), (9, 15), 1, "", "")]),
        ("Venues", "会場", "E2EFDA", [
            ("Venue booking", "会場の予約", "Kanna", "カンナ", (9, 10), (9, 19), 1, "Main square and the old station", "中央広場と旧駅舎"),
            ("Stage build", "ステージ設営", "Utaha", "ウタハ", (9, 21), (10, 16), 0.8, "Backdrops painted at Wildhunt", "背景画はワイルドハントが制作"),
            ("Lighting", "照明", "Hibiki", "ヒビキ", (10, 5), (10, 23), 0.6, "", ""),
            ("Seating", "客席", "Kotori", "コトリ", (10, 12), (10, 23), 0.3, "", "")]),
        ("Booths", "出し物", "FCE4D6", [
            ("Game booth", "ゲームブース", "Momoi", "モモイ", (9, 21), (10, 16), 0.5, "Demo of the new game", "新作ゲームの体験版"),
            ("Food stalls", "屋台", "Fuuka", "フウカ", (9, 14), (10, 9), 0.4, "Haruna wants a tasting menu", "ハルナが試食メニューを希望"),
            ("Sweets café", "スイーツカフェ", "Kazusa", "カズサ", (10, 1), (10, 23), 0.3, "", ""),
            ("Ramen stand", "ラーメン屋台", "Serika", "セリカ", (10, 5), (10, 30), 0.2, "", ""),
            ("Tea party", "お茶会", "Nagisa", "ナギサ", (10, 12), (10, 30), 0, "", ""),
            ("Ninja show", "忍術ショー", "Izuna", "イズナ", (10, 19), (11, 6), 0, "", "")]),
        ("Security", "警備", "FFF2CC", [
            ("Patrol plan", "巡回計画", "Hina", "ヒナ", (9, 7), (9, 18), 1, "", ""),
            ("Crowd control", "誘導", "Kanna", "カンナ", (10, 1), (10, 30), 0.5, "Ferries from Odyssey every 30 min", "オデュッセイアから連絡船を運行"),
            ("First aid", "救護", "Mine", "ミネ", (10, 19), (11, 6), 0, "", ""),
            ("Lost and found", "落とし物", "Kirino", "キリノ", (11, 2), (11, 13), 0, "", "")]),
        ("Opening", "開幕", "EDE1F5", [
            ("Fireworks", "花火", "Aru", "アル", (10, 19), (11, 6), 0, "Waiting on a permit", "許可待ち"),
            ("Rehearsal", "リハーサル", "Sensei", "先生", (11, 9), (11, 13), 0, "", ""),
            ("Opening day", "開幕", "All", "全員", (11, 14), (11, 14), 0, "", "")]),
    ]
    r = 3
    for en, ja, color, tasks in phases:
        ws.cell(row=r, column=1, value=t(en, ja))
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        ws.cell(row=r, column=1).font = Font(bold=True, color="203864")
        ws.cell(row=r, column=1).fill = fill(color)
        r += 1
        for ten, tja, oen, oja, s, e, done, nen, nja in tasks:
            st = "done" if done == 1 else ("blocked" if nen.startswith("Waiting") else ("doing" if done > 0 else "todo"))
            ws.append([t(ten, tja), t(oen, oja), dt.date(2026, *s), dt.date(2026, *e), f"=D{r}-C{r}+1", done,
                       None, t(nen, nja)])
            ws[f"C{r}"].number_format = DAY
            ws[f"D{r}"].number_format = DAY
            ws[f"E{r}"].alignment = Alignment(horizontal="center")
            ws[f"F{r}"].number_format = "0%"
            status_cell(ws[f"G{r}"], st)
            ws[f"H{r}"].font = Font(italic=True, color="595959")
            for c in range(1, 9):
                ws.cell(row=r, column=c).border = box
            r += 1
    ws.cell(row=r, column=1, value=t("Overall progress", "全体の進捗")).font = Font(bold=True)
    ws.cell(row=r, column=6, value=f"=AVERAGE(F4:F{r - 1})").number_format = "0%"
    ws.cell(row=r, column=6).font = Font(bold=True)
    for c in range(1, 9):
        ws.cell(row=r, column=c).border = Border(top=med)
    widths(ws, {"A": 24, "B": 10, "C": 9, "D": 9, "E": 8, "F": 9, "G": 15, "H": 40})
    placeholder_sheets(wb, t("Schools", "学園"), t("Budget", "予算"), t("Notes", "メモ"))
    save(wb, t("Festival Plan.xlsx", "合同祭計画.xlsx"))


# MARK: - Make-up Work Club grades

def grades_iphone():
    wb = Workbook()
    ws = wb.active
    ws.title = t("Mock Exams", "模擬試験")
    cols = [t("Student", "生徒"), t("Test 1", "第1回"), t("Test 2", "第2回"), t("Test 3", "第3回"),
            t("Final", "本試験"), t("Avg", "平均"), t("Grade", "評価")]
    ws.append(cols)
    header(ws, 1, len(cols), "548235")
    students = [("Hifumi", "ヒフミ", 62, 70, 78, 88), ("Azusa", "アズサ", 41, 58, 74, 91),
                ("Koharu", "コハル", 12, 24, 35, 61), ("Hanako", "ハナコ", 3, 0, 100, 100),
                ("Mika", "ミカ", 55, 48, 60, 71), ("Nagisa", "ナギサ", 96, 98, 97, 99),
                ("Seia", "セイア", 99, 100, 98, 100), ("Hasumi", "ハスミ", 84, 86, 88, 90),
                ("Tsurugi", "ツルギ", 72, 75, 70, 78), ("Mashiro", "マシロ", 80, 77, 83, 85)]
    for i, (en, ja, *s) in enumerate(students, start=2):
        ws.append([t(en, ja), *s, f"=AVERAGE(B{i}:E{i})",
                   f'=IFS(F{i}>=90,"A",F{i}>=80,"B",F{i}>=70,"C",F{i}>=60,"D",TRUE,"F")'])
        ws[f"F{i}"].number_format = "0.0"
        ws[f"F{i}"].font = Font(bold=True)
        for c in "BCDEFG":
            ws[f"{c}{i}"].alignment = Alignment(horizontal="center")
    n = 1 + len(students)
    ws.append([t("Average", "平均"), *[f"=AVERAGE({c}2:{c}{n})" for c in "BCDEF"], ""])
    for c in "ABCDEFG":
        cell = ws[f"{c}{n + 1}"]
        cell.font = Font(bold=True, italic=True)
        cell.fill = fill("E2EFDA")
        cell.border = Border(top=Side(style="medium", color="548235"))
        if c in "BCDEF":
            cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="center")
    grid(ws, 2, n, 1, 7)
    widths(ws, {"A": 16, "B": 9, "C": 9, "D": 9, "E": 9, "F": 8, "G": 9})
    save(wb, t("Make-up Work Club.xlsx", "補習授業部.xlsx"))


STUDENTS = [("Hoshino", "ホシノ", "Abydos", 74), ("Shiroko", "シロコ", "Abydos", 82), ("Serika", "セリカ", "Abydos", 70),
            ("Nonomi", "ノノミ", "Abydos", 85), ("Ayane", "アヤネ", "Abydos", 90), ("Yuuka", "ユウカ", "Millennium", 95),
            ("Noa", "ノア", "Millennium", 93), ("Aris", "アリス", "Millennium", 68), ("Momoi", "モモイ", "Millennium", 63),
            ("Midori", "ミドリ", "Millennium", 79), ("Yuzu", "ユズ", "Millennium", 81), ("Hina", "ヒナ", "Gehenna", 96),
            ("Ako", "アコ", "Gehenna", 88), ("Iori", "イオリ", "Gehenna", 72), ("Chinatsu", "チナツ", "Gehenna", 86),
            ("Haruna", "ハルナ", "Gehenna", 84), ("Aru", "アル", "Gehenna", 66), ("Mutsuki", "ムツキ", "Gehenna", 71),
            ("Kayoko", "カヨコ", "Gehenna", 83), ("Haruka", "ハルカ", "Gehenna", 60), ("Mika", "ミカ", "Trinity", 64),
            ("Nagisa", "ナギサ", "Trinity", 94), ("Seia", "セイア", "Trinity", 97), ("Hifumi", "ヒフミ", "Trinity", 77),
            ("Azusa", "アズサ", "Trinity", 75), ("Koharu", "コハル", "Trinity", 52), ("Hanako", "ハナコ", "Trinity", 92),
            ("Tsurugi", "ツルギ", "Trinity", 73), ("Miyako", "ミヤコ", "SRT", 89), ("Kanna", "カンナ", "Valkyrie", 87),
            ("Izuna", "イズナ", "Hyakkiyako", 69), ("Saori", "サオリ", "Arius", 80)]


def grades_ipad():
    wb = Workbook()
    ws = wb.active
    ws.title = t("Term 2", "2学期")
    subjects = [t("Math", "数学"), t("Science", "理科"), t("History", "歴史"), t("Language", "国語"),
                t("English", "英語"), t("Tactics", "戦術"), t("P.E.", "体育"), t("Midterm", "中間"), t("Final", "期末")]
    cols = [t("Student", "生徒"), t("School", "学園"), *subjects, t("Average", "平均"), t("Grade", "評価"),
            t("Rank", "順位")]
    ws.append(cols)
    header(ws, 1, len(cols), "548235")
    n = 1 + len(STUDENTS)
    for i, (en, ja, key, skill) in enumerate(STUDENTS, start=2):
        scores = [max(20, min(100, round(skill + rnd.uniform(-9, 9)))) for _ in subjects]
        ws.append([t(en, ja), school(key), *scores, f"=AVERAGE(C{i}:K{i})",
                   f'=IFS(L{i}>=90,"A",L{i}>=80,"B",L{i}>=70,"C",L{i}>=60,"D",TRUE,"F")',
                   f'=COUNTIF(L$2:L${n},">"&L{i})+1'])
        ws[f"B{i}"].fill = fill(SCHOOL_TINTS[key])
        ws[f"L{i}"].number_format = "0.0"
        ws[f"L{i}"].font = Font(bold=True)
        for c in "CDEFGHIJKLMN":
            ws[f"{c}{i}"].alignment = Alignment(horizontal="center")
    ws.append([t("Average", "平均"), "", *[f"=AVERAGE({c}2:{c}{n})" for c in "CDEFGHIJKL"], "", ""])
    for c in "ABCDEFGHIJKLMN":
        cell = ws[f"{c}{n + 1}"]
        cell.font = Font(bold=True, italic=True)
        cell.fill = fill("E2EFDA")
        cell.border = Border(top=Side(style="medium", color="548235"))
        if c in "CDEFGHIJKL":
            cell.number_format = "0.0"
            cell.alignment = Alignment(horizontal="center")
    grid(ws, 2, n, 1, 14)
    widths(ws, {"A": 12, "B": 14, **{c: 11 for c in "CDEFGHIJK"}, "L": 11, "M": 9, "N": 8})
    placeholder_sheets(wb, t("Term 1", "1学期"), t("Attendance", "出欠"))
    save(wb, t("Supplementary Lessons.xlsx", "補習授業.xlsx"))


# MARK: - Club budgets

CLUBS = [("Foreclosure Task Force", "対策委員会", "Abydos"), ("Seminar", "セミナー", "Millennium"),
         ("Game Dev. Dept.", "ゲーム開発部", "Millennium"), ("Engineering Dept.", "エンジニア部", "Millennium"),
         ("Veritas", "ヴェリタス", "Millennium"), ("Cleaning & Clearing", "C&C", "Millennium"),
         ("Prefect Team", "風紀委員会", "Gehenna"), ("Pandemonium Society", "万魔殿", "Gehenna"),
         ("Gourmet Research", "美食研究会", "Gehenna"), ("Problem Solver 68", "便利屋68", "Gehenna"),
         ("School Lunch Club", "給食部", "Gehenna"), ("Emergency Medicine", "救急医学部", "Gehenna"),
         ("Hot Springs Dev. Dept.", "温泉開発部", "Gehenna"), ("Tea Party", "ティーパーティー", "Trinity"),
         ("Justice Task Force", "正義実現委員会", "Trinity"), ("Sisterhood", "シスターフッド", "Trinity"),
         ("Make-up Work Club", "補習授業部", "Trinity"), ("After-School Sweets", "放課後スイーツ部", "Trinity"),
         ("Remedial Knights", "救護騎士団", "Trinity"), ("RABBIT Squad", "RABBIT小隊", "SRT"),
         ("Public Safety Bureau", "生活安全局", "Valkyrie"), ("Ninjutsu Research Club", "忍術研究部", "Hyakkiyako"),
         ("Yin-Yang Club", "陰陽部", "Hyakkiyako"), ("Festival Committee", "お祭り運営委員会", "Hyakkiyako"),
         ("Arius Squad", "アリウススクワッド", "Arius"), ("Student Council", "生徒会", "Wildhunt"),
         ("Student Council", "生徒会", "Odyssey")]


def clubs_ipad():
    wb = Workbook()
    ws = wb.active
    ws.title = t("Monthly", "月別")
    months = MONTHS[3:12]  # The school year starts in April.
    cols = [t("Club", "部活"), t("School", "学園"), *months, t("Total", "合計"), t("vs Plan", "計画比")]
    ws.append(cols)
    header(ws, 1, len(cols), "C55A11")
    plans = []
    for i, (en, ja, key) in enumerate(CLUBS, start=2):
        base = rnd.randint(18, 40) * 1000
        spend = [round(base * (1 + rnd.uniform(-0.15, 0.2)), -1) for _ in months]
        plans.append(base * len(months))
        ws.append([t(en, ja), school(key), *spend, f"=SUM(C{i}:K{i})", f"=L{i}/{base * len(months)}"])
        ws[f"B{i}"].fill = fill(SCHOOL_TINTS[key])
        for c in "CDEFGHIJKL":
            ws[f"{c}{i}"].number_format = CREDITS
        ws[f"L{i}"].font = Font(bold=True)
        ws[f"M{i}"].number_format = "0.0%"
    n = 1 + len(CLUBS)
    ws.append([t("Total", "合計"), "", *[f"=SUM({c}2:{c}{n})" for c in "CDEFGHIJKL"], f"=L{n + 1}/{sum(plans)}"])
    for c in "ABCDEFGHIJKLM":
        cell = ws[f"{c}{n + 1}"]
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill("843C0C")
        cell.number_format = "0.0%" if c == "M" else CREDITS
    grid(ws, 2, n, 1, 13)
    widths(ws, {"A": 26, "B": 16, **{c: 12 for c in "CDEFGHIJK"}, "L": 13, "M": 11})

    by_school = wb.create_sheet(t("By School", "学園別"))
    by_school.append([t("School", "学園"), t("Total", "合計"), t("Share", "割合")])
    header(by_school, 1, 3, "C55A11")
    keys = list(dict.fromkeys(key for _, _, key in CLUBS))
    for i, key in enumerate(keys, start=2):
        by_school.append([school(key), f"=SUMIFS({ws.title}!L2:L{n},{ws.title}!B2:B{n},A{i})",
                          f"=B{i}/SUM(B$2:B${1 + len(keys)})"])
        by_school[f"B{i}"].number_format = CREDITS
        by_school[f"C{i}"].number_format = "0.0%"
    grid(by_school, 2, 1 + len(keys), 1, 3)
    widths(by_school, {"A": 14, "B": 13, "C": 9})
    placeholder_sheets(wb, t("Requests", "申請"))
    save(wb, t("Club Budgets.xlsx", "部活予算.xlsx"))


# MARK: - Delimited text

def csvs():
    with open(os.path.join(OUT, t("Schale Supplies.csv", "シャーレ備品.csv")), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([t("Item", "品名"), t("Category", "分類"), t("In Stock", "在庫"), t("Reorder At", "発注点")])
        for en, ja, cen, cja, qty, reorder in [
                ("Activity reports", "活動報告書", "Paperwork", "書類", 340, 100),
                ("Energy drinks", "エナジードリンク", "Snacks", "おやつ", 12, 24),
                ("Peroro plushies", "ペロロ様ぬいぐるみ", "Goods", "グッズ", 3, 1),
                ("Tactical training discs", "戦術教育BD", "Training", "教材", 58, 20),
                ("Workbooks", "ワークブック", "Training", "教材", 140, 50),
                ("Eligma", "神名のカケラ", "Rare", "希少", 7, 0),
                ("Instant ramen", "カップラーメン", "Snacks", "おやつ", 36, 30),
                ("Spare halos", "予備のヘイロー", "Rare", "希少", 0, 0)]:
            w.writerow([t(en, ja), t(cen, cja), qty, reorder])
    with open(os.path.join(OUT, t("Patrol Log.csv", "パトロール記録.csv")), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow([t("Date", "日付"), t("District", "地区"), t("Students", "生徒"), t("Incidents", "件数")])
        for d, en, ja, students, incidents in [
                ("2026-10-01", "Abydos desert", "アビドス砂漠", t("Shiroko, Hoshino", "シロコ、ホシノ"), 3),
                ("2026-10-02", "D.U. Shiratori", "D.U.シラトリ区", t("Kanna, Kirino", "カンナ、キリノ"), 5),
                ("2026-10-03", "Black Market", "ブラックマーケット", t("Aru, Kayoko", "アル、カヨコ"), 9),
                ("2026-10-04", "Trinity Square", "トリニティ広場", t("Tsurugi, Hasumi", "ツルギ、ハスミ"), 1),
                ("2026-10-05", "Odyssey harbour", "オデュッセイア港", t("Miyako, Saki", "ミヤコ、サキ"), 2)]:
            w.writerow([d, t(en, ja), students, incidents])


if DEVICE == "iPhone":
    budget_iphone()
    grades_iphone()
    plan_iphone()
else:
    budget_ipad()
    clubs_ipad()
    plan_ipad()
    grades_ipad()
csvs()
print("\n".join(sorted(os.listdir(OUT))))
